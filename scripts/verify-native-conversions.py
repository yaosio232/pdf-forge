"""Native conversion regression harness.

Requires Windows desktop Word/WebView2 and Python reportlab, python-docx, Pillow,
pypdf, and lxml. Never manipulates application UI. Artifacts remain under the
chosen output folder. Every worker has a deadline; cleanup targets only its PID
process tree and Word PID plus creation-time marker. Use --long-paths for genuine
>260-character source/output paths and --edges for native renderer failures plus
existing Word instance safety. This is direct function verification, not UI QA.
"""
import argparse, ctypes, hashlib, json, os, pathlib, subprocess, time
from ctypes import wintypes
from reportlab.pdfgen import canvas
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader
from worker_resources import runtime_directory

ROOT=pathlib.Path(__file__).resolve().parent.parent
parser=argparse.ArgumentParser(description='Real Windows worker conversion checks; synthetic documents only, no UI operations.')
parser.add_argument('--exe',type=pathlib.Path,default=ROOT/'tauri/src-tauri/target/debug/pdf-forge.exe')
parser.add_argument('--output-dir',type=pathlib.Path,default=ROOT/'TestResults/functional-conversions')
parser.add_argument('--label',default='native')
parser.add_argument('--long-paths',action='store_true')
parser.add_argument('--edges',action='store_true',help='Keep an owned existing Word instance alive during conversion and test actual formula/diagram failures.')
parser.add_argument('--timeout',type=int,default=65)
parser.add_argument('--markdown-assets',type=pathlib.Path,help='Defaults to the actual executable resources/markdown directory.')
cli=parser.parse_args()
BASE=cli.output_dir.resolve(); BASE.mkdir(parents=True,exist_ok=True)
EXE=cli.exe.resolve()
if not EXE.is_file(): parser.error('Worker executable does not exist')
ASSETS=(cli.markdown_assets or runtime_directory(EXE,BASE)/'resources/markdown').resolve()
if not all((ASSETS/name).is_file() for name in ['renderer.js','renderer.css']): parser.error('Offline renderer assets are missing; specify --markdown-assets for a development worker')
with EXE.open('rb') as file: executable_hash=hashlib.file_digest(file,'sha256').hexdigest()
(BASE/('run-config-'+cli.label+'.json')).write_text(json.dumps({'executable':str(EXE),'executableSha256':executable_hash,'markdownAssets':str(ASSETS),'longPaths':cli.long_paths,'timeoutSeconds':cli.timeout},indent=2),encoding='utf-8')
long_base=BASE/'long-paths'/('a'*90)/('b'*90)
fixtures=(long_base if cli.long_paths else BASE)/'fixtures'; fixtures.mkdir(exist_ok=True,parents=True)
pdf=fixtures/'synthetic.pdf'; c=canvas.Canvas(str(pdf))
doc=Document()
for n in range(1,21):
    c.setFont('Helvetica-Bold',18); c.drawString(50,790,f'FORGE PAGE {n:02d}')
    c.setFont('Helvetica',12)
    for row in range(12): c.drawString(50,750-row*25,f'Synthetic paragraph {row+1}: page {n}, conversion verification data.')
    c.showPage()
    doc.add_heading(f'FORGE PAGE {n:02d}',level=1)
    for row in range(8): doc.add_paragraph(f'Synthetic paragraph {row+1}: page {n}, conversion verification data.')
    table=doc.add_table(rows=2,cols=2); table.cell(0,0).text='Page'; table.cell(0,1).text=str(n); table.cell(1,0).text='Value'; table.cell(1,1).text=str(n*7)
    if n<20: doc.add_page_break()
c.save(); doc.save(fixtures/'synthetic.docx')
im=Image.new('RGB',(320,160),'white'); ImageDraw.Draw(im).text((20,70),'FORGE LOCAL IMAGE',fill='black'); im.save(fixtures/'local.png')
(fixtures/'simple.md').write_text('\n\n'.join(f'# FORGE SECTION {n:02d}\n\n'+('Synthetic Markdown paragraph with **bold** and a table. '*70)+'\n\n|Page|Value|\n|---|---|\n|'+str(n)+'|'+str(n*7)+'|' for n in range(1,21)),encoding='utf-8')
(fixtures/'advanced.md').write_text('# FORGE ADVANCED\n\n![Local fixture](local.png)\n\n$$E=mc^2$$\n\nInline math $x^2+y^2=z^2$.\n\n```mermaid\nflowchart LR\n A[Source] --> B[Output]\n```\n\n'+(fixtures/'simple.md').read_text(encoding='utf-8'),encoding='utf-8')

def cleanup_word(marker):
    if not marker.exists(): return
    data=json.loads(marker.read_text(encoding='utf-8-sig')); k=ctypes.WinDLL('kernel32',use_last_error=True)
    k.OpenProcess.restype=wintypes.HANDLE; k.OpenProcess.argtypes=[wintypes.DWORD,wintypes.BOOL,wintypes.DWORD]
    handle=k.OpenProcess(0x1000|1,False,data['pid'])
    if not handle:return
    k.GetProcessTimes.argtypes=[wintypes.HANDLE]+[ctypes.POINTER(wintypes.FILETIME)]*4
    k.TerminateProcess.argtypes=[wintypes.HANDLE,wintypes.UINT]; k.CloseHandle.argtypes=[wintypes.HANDLE]
    vals=[wintypes.FILETIME() for _ in range(4)]
    if k.GetProcessTimes(handle,*[ctypes.byref(v) for v in vals]):
        stamp=(vals[0].dwHighDateTime<<32)|vals[0].dwLowDateTime
        if stamp==data['creationTime']:k.TerminateProcess(handle,1)
    k.CloseHandle(handle)

results=[]
for name,kind,source,ext,opts in [
 ('word-pdf','wordToPdf','synthetic.docx','pdf',None),
 ('pdf-word','pdfToWord','synthetic.pdf','docx',None),
 ('markdown-simple','markdown','simple.md','pdf',dict(toc=False,header='',footer='',pageNumbers=False)),
 ('markdown-advanced','markdown','advanced.md','pdf',dict(toc=True,header='FORGE HEADER',footer='FORGE FOOTER',pageNumbers=True))]:
    run=BASE/(name+'-'+cli.label);run.mkdir(exist_ok=True)
    for marker in ['cancel','result.json','word.json']:
        if (run/marker).exists(): (run/marker).unlink()
    output_parent=(long_base/'outputs' if cli.long_paths else BASE)/run.name;output_parent.mkdir(exist_ok=True,parents=True)
    output=output_parent/('output.'+ext)
    if output.exists():output.unlink()
    source_path=fixtures/source;source_hash=hashlib.sha256(source_path.read_bytes()).hexdigest()
    prefix='\\\\?\\' if cli.long_paths else ''
    request=dict(kind=kind,source=prefix+str(source_path),output=prefix+str(output))
    if opts is not None:request['options']=opts
    (run/'request.json').write_text(json.dumps(request),encoding='utf-8')
    env=os.environ.copy();env.update(PDF_FORGE_WORD_IDENTITY=str(run/'word.json'),PDF_FORGE_CANCEL_FILE=str(run/'cancel'),PDF_FORGE_MARKDOWN_ASSETS=str(ASSETS))
    start=time.monotonic(); p=subprocess.Popen([str(EXE),'--worker',str(run)],env=env,stdout=subprocess.PIPE,stderr=subprocess.PIPE,creationflags=subprocess.CREATE_NO_WINDOW)
    timedout=False
    try:out,err=p.communicate(timeout=cli.timeout)
    except subprocess.TimeoutExpired:
        timedout=True;(run/'cancel').touch()
        try:out,err=p.communicate(timeout=5)
        except subprocess.TimeoutExpired:
            subprocess.run(['taskkill','/PID',str(p.pid),'/T','/F'],capture_output=True,creationflags=subprocess.CREATE_NO_WINDOW);out,err=p.communicate(timeout=10)
        cleanup_word(run/'word.json')
    entry=dict(name=name,seconds=round(time.monotonic()-start,2),exit=p.returncode,timeout=timedout,sourceHashBefore=source_hash,sourceHashAfter=hashlib.sha256(source_path.read_bytes()).hexdigest(),stdout=out.decode(errors='replace'),stderr=err.decode(errors='replace'))
    if (run/'result.json').exists():entry['response']=json.loads((run/'result.json').read_text(encoding='utf-8'))
    if output.exists():
        entry['bytes']=output.stat().st_size
        if ext=='pdf':
            try:
                r=PdfReader(output);text='\n'.join(page.extract_text() or '' for page in r.pages);entry.update(pages=len(r.pages),textSample=text[:400],containsFirst='FORGE' in text,containsLast='20' in text,header='FORGE HEADER' in text,footer='FORGE FOOTER' in text)
            except Exception as e:entry['inspectionError']=str(e)
        else:
            try:d=Document(output);entry.update(paragraphs=len(d.paragraphs),textSample='\n'.join(x.text for x in d.paragraphs)[:400])
            except Exception as e:entry['inspectionError']=str(e)
    results.append(entry);print(json.dumps({key:entry[key] for key in ['name','seconds','response','sourceHashBefore','sourceHashAfter'] if key in entry}),flush=True)
    (BASE/('results-'+cli.label+'.json')).write_text(json.dumps(results,indent=2),encoding='utf-8')



from inspect_native_conversions import inspect_outputs
inspection=inspect_outputs(BASE,cli.label)
if not all(entry["passed"] for entry in inspection): raise SystemExit(1)

if cli.edges:
    base=BASE; root=ROOT; exe=EXE
    (base/"hold-owned-word.ps1").write_text('param([string]$Directory)\n$ErrorActionPreference=\'Stop\'\nAdd-Type \'using System; using System.Runtime.InteropServices; public static class ForgeWindowPid { [DllImport("user32.dll")] public static extern uint GetWindowThreadProcessId(IntPtr hwnd, out uint pid); }\'\n$word=$null; $document=$null\ntry {\n  $word=New-Object -ComObject Word.Application\n  $word.Visible=$false; $word.DisplayAlerts=0\n  $document=$word.Documents.Add()\n  $document.Content.Text=\'FORGE EXISTING OWNED WORD SENTINEL\'\n  [uint32]$wordPid=0\n  [void][ForgeWindowPid]::GetWindowThreadProcessId([IntPtr]$document.Windows.Item(1).Hwnd,[ref]$wordPid)\n  $process=Get-Process -Id $wordPid\n  @{pid=$wordPid;creationTime=$process.StartTime.ToUniversalTime().ToFileTimeUtc();path=$process.Path} | ConvertTo-Json | Set-Content -Encoding utf8 (Join-Path $Directory \'existing.json\')\n  $deadline=(Get-Date).AddSeconds(55)\n  while (!(Test-Path (Join-Path $Directory \'release\')) -and (Get-Date) -lt $deadline) { Start-Sleep -Milliseconds 100 }\n  @{sentinel=$document.Content.Text;pid=$wordPid;alive=!(Get-Process -Id $wordPid).HasExited} | ConvertTo-Json | Set-Content -Encoding utf8 (Join-Path $Directory \'existing-after.json\')\n} finally {\n  if ($document) { $document.Close(0) }\n  if ($word) { $word.Quit(0) }\n}\n',encoding="utf-8")
    reports=[]
    def invoke(name,task):
        directory=base/(name+'-'+cli.label);directory.mkdir(exist_ok=True)
        for marker in ['cancel','result.json','word.json']:
            if (directory/marker).exists(): (directory/marker).unlink()
        (directory/'request.json').write_text(json.dumps(task),encoding='utf-8')
        env=os.environ.copy();env.update(PDF_FORGE_WORD_IDENTITY=str(directory/'word.json'),PDF_FORGE_CANCEL_FILE=str(directory/'cancel'),PDF_FORGE_MARKDOWN_ASSETS=str(ASSETS))
        p=subprocess.Popen([str(exe),'--worker',str(directory)],env=env,stdout=subprocess.PIPE,stderr=subprocess.PIPE,creationflags=subprocess.CREATE_NO_WINDOW)
        start=time.monotonic()
        try:out,err=p.communicate(timeout=cli.timeout)
        except subprocess.TimeoutExpired:
            (directory/'cancel').touch()
            try:out,err=p.communicate(timeout=5)
            except subprocess.TimeoutExpired:
                subprocess.run(['taskkill','/PID',str(p.pid),'/T','/F'],capture_output=True,creationflags=subprocess.CREATE_NO_WINDOW);out,err=p.communicate(timeout=10)
            cleanup_word(directory/'word.json')
            raise RuntimeError('Native edge worker timed out')
        response=json.loads((directory/'result.json').read_text(encoding='utf-8'))
        report=dict(name=name,seconds=round(time.monotonic()-start,2),response=response);reports.append(report);print(json.dumps(report),flush=True)
        return directory,response

    holder=base/('existing-word-'+cli.label);holder.mkdir(exist_ok=True)
    for n in ['release','existing.json','existing-after.json']:
        if (holder/n).exists():(holder/n).unlink()
    ps=subprocess.Popen(['powershell','-NoProfile','-ExecutionPolicy','Bypass','-File',str(base/'hold-owned-word.ps1'),'-Directory',str(holder)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,creationflags=subprocess.CREATE_NO_WINDOW)
    try:
        deadline=time.monotonic()+15
        while not (holder/'existing.json').exists():
            if ps.poll() is not None:raise RuntimeError(ps.communicate()[1].decode(errors='replace'))
            if time.monotonic()>deadline:raise RuntimeError('Word holder startup timeout')
            time.sleep(.1)
        original=json.loads((holder/'existing.json').read_text(encoding='utf-8-sig'))
        output=base/('existing-safe-'+cli.label)/'output.pdf';output.parent.mkdir(exist_ok=True)
        if output.exists():output.unlink()
        directory,response=invoke('existing-safe',dict(kind='wordToPdf',source=str(fixtures/'synthetic.docx'),output=str(output)))
    finally:
        (holder/'release').touch()
        try: out,err=ps.communicate(timeout=10)
        except subprocess.TimeoutExpired:
            subprocess.run(['taskkill','/PID',str(ps.pid),'/T','/F'],capture_output=True,creationflags=subprocess.CREATE_NO_WINDOW)
            out,err=ps.communicate(timeout=10)
        finally: cleanup_word(holder/'existing.json')
    after=json.loads((holder/'existing-after.json').read_text(encoding='utf-8-sig'))
    worker=json.loads((directory/'word.json').read_text(encoding='utf-8')) if (directory/'word.json').exists() else None
    reports[-1].update(existingIdentity=original,existingAfter=after,workerIdentity=worker,passed=after['alive'] and 'FORGE EXISTING OWNED WORD SENTINEL' in after['sentinel'] and (worker is None or worker['pid']!=original['pid']))

    options=dict(toc=False,header='',footer='',pageNumbers=False)
    for name,content,code in [('invalid-math','# Formula\n\n$$\\notARealCommand{x}$$','MARKDOWN_MATH'),('invalid-mermaid','# Diagram\n\n```mermaid\nflowchart LR\n A --> [\n```','MARKDOWN_MERMAID')]:
        source=fixtures/f'{name}.md';source.write_text(content,encoding='utf-8');before=hashlib.sha256(source.read_bytes()).hexdigest()
        output=base/(name+'-'+cli.label)/'output.pdf'
        if output.exists():output.unlink()
        _,response=invoke(name,dict(kind='markdown',source=str(source),output=str(output),options=options))
        reports[-1].update(passed=response['error'] is not None and response['error']['code']==code and not output.exists(),sourceUnchanged=before==hashlib.sha256(source.read_bytes()).hexdigest())
    (base/('native-edges-'+cli.label+'.json')).write_text(json.dumps(reports,indent=2),encoding='utf-8');print(json.dumps(reports,indent=2))

    if not all(entry.get("passed") and entry.get("sourceUnchanged",True) for entry in reports): raise SystemExit(1)

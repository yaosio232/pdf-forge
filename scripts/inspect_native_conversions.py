import hashlib,json,pathlib,re,sys,zipfile
from lxml import etree
from pypdf import PdfReader
from docx import Document

def inspect_outputs(base: pathlib.Path, label: str):
    executions={item['name']:item for item in json.loads((base/('results-'+label+'.json')).read_text(encoding='utf-8'))}
    report=[]
    for name in ['word-pdf','pdf-word','markdown-simple','markdown-advanced']:
        directory=base/(name+'-'+label)
        response=json.loads((directory/'result.json').read_text(encoding='utf-8'))
        ext='docx' if name=='pdf-word' else 'pdf';request=json.loads((directory/'request.json').read_text(encoding='utf-8'));path=pathlib.Path(request['output'])
        result={'name':name,'workerResponse':response,'checks':{}}
        execution=executions[name]
        result['checks']['sourceHashUnchanged']=execution['sourceHashBefore']==execution['sourceHashAfter']
        result['checks']['workerCompletedWithinDeadline']=not execution['timeout'] and execution['exit']==0
        result['sourcePathCharacters']=len(request['source']);result['outputPathCharacters']=len(request['output'])
        if path.exists():
            result['sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
            if ext=='pdf':
                pdf=PdfReader(path);texts=[page.extract_text() or '' for page in pdf.pages];text='\n'.join(texts)
                result['pages']=len(pdf.pages)
                result['checks']['richDocumentAtLeast20Pages']=len(pdf.pages)>=20
                ids=[f'FORGE PAGE {n:02d}' if name=='word-pdf' else f'FORGE SECTION {n:02d}' for n in range(1,21)]
                result['checks']['all20DistinctIds']=all(i in text for i in ids)
                result['checks']['pageIdOrder']=all(text.find(ids[n])<text.find(ids[n+1]) for n in range(19))
                result['checks']['a4']=all(abs(float(p.mediabox.width)-595.28)<2 and abs(float(p.mediabox.height)-841.89)<2 for p in pdf.pages) if name.startswith('markdown') else True
                if name=='markdown-advanced':
                    result['checks']['toc']=('Contents' in texts[0] and 'FORGE SECTION 20' in texts[0])
                    result['checks']['headerEveryPage']=all('FORGE HEADER' in t for t in texts)
                    result['checks']['footerEveryPage']=all('FORGE FOOTER' in t for t in texts)
                    result['checks']['pageCounterEveryPage']=all(re.search(rf'{i+1}\s*/\s*{len(texts)}',t) for i,t in enumerate(texts))
                    result['checks']['mermaidText']='Source' in text and 'Output' in text
                    result['checks']['mathText']='mc' in text and 'Inline math' in text
                    result['checks']['embeddedImage']=any(list(page.images) for page in pdf.pages)
                    start=text.rfind('FORGE ADVANCED')
                    result['mathAndDiagramText']=text[start:start+400]
            else:
                doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
                result['checks']['all20DistinctIds']=all(f'FORGE PAGE {n:02d}' in text for n in range(1,21))
                with zipfile.ZipFile(path) as archive:
                    parts=[n for n in archive.namelist() if n.startswith('word/') and n.endswith('.xml')]
                    all_text='\n'.join(''.join(etree.fromstring(archive.read(n)).itertext()) for n in parts)
                result['checks']['richParagraphsIncludingHeadersFooters']=all(f'Synthetic paragraph {n}:' in all_text for n in range(1,13))
            result['checks']['workerSuccess']=response['error'] is None
        else:result['checks']['outputExists']=False
        result['passed']=all(bool(v) for v in result['checks'].values());report.append(result)
    (base/('inspection-'+label+'.json')).write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
    print(json.dumps([{'name':entry['name'],'passed':entry['passed'],'pages':entry.get('pages'),'failedChecks':[key for key,value in entry['checks'].items() if not value]} for entry in report]))
    return report

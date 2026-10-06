"""Direct, real supervisor workflows and independent output checks; no UI operation."""
import argparse,hashlib,json,os,pathlib,shutil,subprocess
from pypdf import PdfReader
from docx import Document
from worker_resources import runtime_directory
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def main():
    parser=argparse.ArgumentParser();parser.add_argument('directory',type=pathlib.Path);parser.add_argument('runner',type=pathlib.Path);parser.add_argument('worker',type=pathlib.Path);parser.add_argument('fixtures',type=pathlib.Path)
    args=parser.parse_args();root=args.directory.resolve();root.mkdir(parents=True,exist_ok=False)
    fixture_dir=root/'inputs';shutil.copytree(args.fixtures,fixture_dir)
    simple=fixture_dir/'simple.md';simple.write_text('# FORGE SIMPLE\n\n'+simple.read_text(encoding='utf-8'),encoding='utf-8')
    word=Document(fixture_dir/'synthetic.docx');word.paragraphs[0].text='FORGE WORD PAGE 01';word.save(fixture_dir/'synthetic.docx')
    shutil.copyfile(fixture_dir/'synthetic.pdf',fixture_dir/'pdf-conversion-source.pdf')
    source_hashes={str(path):sha(path) for path in fixture_dir.iterdir() if path.is_file()}
    runtime=runtime_directory(args.worker,root);renderer=runtime/'resources'/'markdown';assert renderer.is_dir(),renderer
    env={**os.environ,'PATH':str(runtime)+os.pathsep+os.environ['PATH'],'PDF_FORGE_MARKDOWN_ASSETS':str(renderer)}
    results=[]
    def run(name,request,cancel=False):
        input=root/f'{name}-request.json';output=root/f'{name}-result.json';input.write_text(json.dumps(request),encoding='utf-8')
        cmd=[str(args.runner.resolve()),str(input),str(args.worker.resolve()),str(output),'--auto']
        if cancel:cmd.append('--cancel-after-worker-start')
        process=subprocess.run(cmd,capture_output=True,text=True,env=env,timeout=420)
        assert process.returncode==0,(name,process.returncode,process.stderr)
        event=json.loads(output.read_text(encoding='utf-8'));results.append({'id':name,'event':event,'workerExecutableSha256':sha(args.worker.resolve())});return event
    options=dict(toc=True,header='WORKFLOW HEADER',footer='WORKFLOW FOOTER',pageNumbers=True)
    simple=fixture_dir/'simple.md';advanced=fixture_dir/'advanced.md';invalid=fixture_dir/'invalid.md';invalid.write_text('# Invalid diagram\n\n```mermaid\nflowchart LR\n A --> [broken\n```\n',encoding='utf-8')
    request=lambda sources,merged,overwrite=False:dict(kind='markdown',sources=[str(p) for p in sources],output='',merged=merged,overwrite=overwrite,options=options)
    event=run('MD-INDIVIDUAL-PARTIAL',request([simple,invalid,advanced],False));assert event['status']=='partial',event
    assert [i['error'] is None for i in event['items']]==[True,False,True]
    assert event['items'][1]['error']['code']=='MARKDOWN_MERMAID'
    for item in [event['items'][0],event['items'][2]]:
        path=pathlib.Path(item['output']);pdf=PdfReader(path);assert len(pdf.pages)>=20
        text='\n'.join(page.extract_text() for page in pdf.pages);assert 'WORKFLOW HEADER' in text and 'FORGE SECTION 20' in text
        assert os.path.samefile(path.parent,fixture_dir);results[-1].setdefault('outputHashes',{})[str(path)]=sha(path)
    final=fixture_dir/'simple_merged.pdf';shutil.copyfile(fixture_dir/'synthetic.pdf',final);previous=sha(final)
    event=run('MD-MERGED-FAILURE-PRESERVES-OUTPUT',request([simple,invalid,advanced],True,True));assert event['status']=='failed' and event['error']['code']=='BATCH_FAILED' and sha(final)==previous,event
    event=run('MD-SORTED-MERGE',request([advanced,simple],True));assert event['status']=='completed',event
    path=pathlib.Path(event['items'][0]['output']);pdf=PdfReader(path);text='\n'.join(page.extract_text() for page in pdf.pages)
    assert text.index('FORGE ADVANCED')<text.index('FORGE SIMPLE')
    assert len(pdf.pages)==len(PdfReader(fixture_dir/'advanced.pdf').pages)+len(PdfReader(fixture_dir/'simple.pdf').pages)
    results[-1]['outputHash']=sha(path)
    badword=fixture_dir/'invalid.docx';badword.write_bytes(b'invalid package')
    event=run('MIXED-WORD-PDF-SORTED-PARTIAL',dict(kind='merge',sources=[str(fixture_dir/'synthetic.docx'),str(badword),str(fixture_dir/'synthetic.pdf')],output='',overwrite=False));assert event['status']=='partial',event
    assert [i['error'] is None for i in event['items']]==[True,False,True]
    path=pathlib.Path(event['items'][0]['output']);pdf=PdfReader(path);assert len(pdf.pages)==40
    text='\n'.join(page.extract_text() for page in pdf.pages);assert 'FORGE WORD PAGE 01' in text and 'FORGE PAGE 20' in text,text[:250]
    assert 'FORGE WORD PAGE 01' in pdf.pages[0].extract_text() and 'FORGE PAGE 01' in pdf.pages[20].extract_text()
    results[-1]['outputHash']=sha(path)
    event=run('PDF-TO-WORD-AUTOMATIC-OUTPUT',dict(kind='pdfToWord',source=str(fixture_dir/'pdf-conversion-source.pdf'),output='',overwrite=False));assert event['status']=='completed',event
    assert os.path.samefile(event['items'][0]['output'],fixture_dir/'pdf-conversion-source.docx')
    results[-1]['outputHash']=sha(pathlib.Path(event['items'][0]['output']))
    document=Document(event['items'][0]['output']);text='\n'.join(p.text for p in document.paragraphs)
    for section in document.sections:
        text+='\n'+'\n'.join(p.text for p in section.header.paragraphs+section.footer.paragraphs)
    assert all(f'FORGE PAGE {n:02d}' in text for n in range(1,21))
    for name,digest in source_hashes.items():
        assert sha(pathlib.Path(name))==digest,name
    (root/'results.json').write_text(json.dumps(results,indent=2),encoding='utf-8');print(json.dumps([{'id':r['id'],'status':r['event']['status']} for r in results],indent=2))
if __name__=='__main__':main()
